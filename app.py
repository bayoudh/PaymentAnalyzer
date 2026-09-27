from flask import Flask, render_template, request, send_file, url_for
import pandas as pd
import io
import re
import uuid
import os
import sys
import threading
import webbrowser


# ============================================================
# CONFIGURATION FOR NORMAL PYTHON + PYINSTALLER EXE
# ============================================================

if getattr(sys, "frozen", False):
    # When running as a PyInstaller EXE:
    # BASE_DIR = folder containing PaymentAnalyzer.exe
    BASE_DIR = os.path.dirname(sys.executable)

    # PyInstaller extracts --add-data files into _MEIPASS
    RESOURCE_DIR = sys._MEIPASS
else:
    # When running app.py normally
    BASE_DIR = os.path.dirname(os.path.abspath(__file__))
    RESOURCE_DIR = BASE_DIR


TEMPLATES_FOLDER = os.path.join(
    RESOURCE_DIR,
    "templates"
)

STATIC_FOLDER = os.path.join(
    RESOURCE_DIR,
    "static"
)

# Generated Excel files must stay OUTSIDE the PyInstaller
# temporary extraction directory, beside the EXE.
GENERATED_FOLDER = os.path.join(
    BASE_DIR,
    "generated"
)

os.makedirs(
    GENERATED_FOLDER,
    exist_ok=True
)


app = Flask(
    __name__,
    template_folder=TEMPLATES_FOLDER,
    static_folder=STATIC_FOLDER
)

# Maximum upload size: 50 MB per request
app.config["MAX_CONTENT_LENGTH"] = 50 * 1024 * 1024



# ============================================================
# NETTOYER NUMERO CARTE
# ============================================================

def clean_card(value):
    if pd.isna(value):
        return ""

    value = str(value).strip().replace("\ufeff", "")

    # Exemple : ="6665767565100" ou ="123" ou = "123"
    value = value.strip()
    if value.startswith("="):
        value = value[1:].strip().strip('"').strip("'").strip()

    if len(value) >= 2 and (
        (value.startswith('"') and value.endswith('"'))
        or (value.startswith("'") and value.endswith("'"))
    ):
        value = value[1:-1].strip()

    # Supprimer espaces (incl. insécables), tirets, apostrophes
    for sep in (" ", "\xa0", "\u202f", "\u2007", "\t", "-", "'", "’", "`"):
        value = value.replace(sep, "")

    # Exemple scientifique :
    # 6,6657675651E+12
    # ou 6.6657675651E+12
    if "E+" in value.upper():
        try:
            value = value.replace(",", ".")

            number = float(value)

            # Convertir en entier
            value = str(int(number))

        except (ValueError, OverflowError):
            pass

    # Exemple : 6665767565100.0
    if re.fullmatch(r"\d+\.0", value):
        value = value[:-2]

    return value.strip()


def card_precision_risk(raw):
    """True si un numéro de carte risque d'être tronqué par Excel.

    Excel ne garde que 15 chiffres significatifs : un .xlsx converti où
    la colonne carte est restée au format Nombre perd les derniers
    chiffres (ex. ...5003 -> ...5056). Seul le CSV d'origine ou une
    colonne formatée en Texte avant conversion garde le numéro exact.
    """
    if pd.isna(raw):
        return False
    text = str(raw).strip().replace(" ", "").replace("\xa0", "")
    if text.startswith("="):
        return False  # format ="..." : texte, pas de perte
    digits = re.sub(r"\D", "", text)
    if len(digits) <= 16:  # 15 chiffres + marge (.0 / exposant)
        return False
    # Flottant / scientifique / décimal .0 sur plus de 15 chiffres
    return bool(
        "E+" in text.upper()
        or re.fullmatch(r"\d+\.0+", text)
        or re.fullmatch(r"\d+\.\d+", text)
    )


# ============================================================
# NETTOYER MONTANT
# ============================================================

def _parse_amount_number(value):
    """Parse un montant vers float. Retourne None si non numérique."""
    if pd.isna(value):
        return None

    text = str(value).strip().replace("\ufeff", "")

    if text == "" or text in ("-", "—", "--", "N/A", "NA", "#N/A"):
        return None

    # Négatif entre parenthèses : (1 266,70) -> -1266.70
    is_negative = False
    if text.startswith("(") and text.endswith(")"):
        is_negative = True
        text = text[1:-1].strip()

    # Trailing minus : 123,45- -> -123.45
    if text.endswith("-") and not text.startswith("-"):
        is_negative = True
        text = text[:-1].strip()

    # Supprimer espaces (incl. insécables) et apostrophes milliers CH/FR
    for sep in (" ", "\xa0", "\u202f", "\u2007", "\t", "'", "’", "`"):
        text = text.replace(sep, "")

    # Supprimer symboles monétaires / lettres (DT, TND, EUR, %, etc.)
    # On garde uniquement chiffres, virgule, point, moins
    text = re.sub(r"[^0-9,.\-]", "", text)

    if text in ("", "-", ".", ",", "-.", ".-"):
        return None

    # Cas avec virgule ET point : le dernier séparateur est la décimale
    # 1.266,70 (EU) -> 1266.70 | 1,266.70 (US) -> 1266.70
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            # EU : points = milliers
            text = text.replace(".", "")
            text = text.replace(",", ".")
        else:
            # US : virgules = milliers
            text = text.replace(",", "")
    elif "," in text:
        # Que des virgules :
        # 866,7 -> 866.7 (décimale EU)
        # 1,266 -> 1266 (milliers US, 3 chiffres)
        # 1,266,700 -> 1266700 (milliers, dernière = décimale si 1-2 chiffres)
        if text.count(",") > 1:
            head, _, tail = text.rpartition(",")
            if len(tail) <= 2:
                head = head.replace(",", "")
                text = head + "." + tail
            else:
                text = text.replace(",", "")
        else:
            head, _, tail = text.rpartition(",")
            if len(tail) == 3 and head.lstrip("-").isdigit():
                text = head + tail  # milliers : 1,266 -> 1266
            else:
                text = head + "." + tail  # décimale : 866,7 -> 866.7
    else:
        # Que des points : 1266.70 -> décimal, 1.266.700 -> milliers
        if text.count(".") > 1:
            _, _, tail = text.rpartition(".")
            if len(tail) == 3:
                text = text.replace(".", "")  # milliers : 1.266.700 -> 1266700
            else:
                head, _, tail = text.rpartition(".")
                head = head.replace(".", "")
                text = head + "." + tail

    try:
        number = float(text)
        if is_negative and number > 0:
            number = -number
        return number
    except ValueError:
        return None


def clean_amount(value):
    number = _parse_amount_number(value)
    return 0.0 if number is None else number


# ============================================================
# NETTOYER DATE -> JJ/MM/AAAA
# ============================================================

def clean_date(value):
    """Normalise une date vers 'JJ/MM/AAAA'. Retourne '' si illisible."""
    if pd.isna(value):
        return ""

    text = str(value).strip().replace("\ufeff", "")
    if text == "":
        return ""

    # Garder la partie date si datetime avec heure : "23/09/2026 06:16"
    text = text.split()[0].strip().replace(".", "/").replace("-", "/")

    # Format AAAA/MM/JJ ou AAAA-MM-JJ -> JJ/MM/AAAA
    m = re.fullmatch(r"(\d{4})/(\d{1,2})/(\d{1,2})", text)
    if m:
        return f"{int(m.group(3)):02d}/{int(m.group(2)):02d}/{m.group(1)}"

    # Format JJ/MM/AAAA (ou J/M/AA)
    m = re.fullmatch(r"(\d{1,2})/(\d{1,2})/(\d{2,4})", text)
    if m:
        year = m.group(3)
        if len(year) == 2:
            year = "20" + year
        try:
            d, mo = int(m.group(1)), int(m.group(2))
            if 1 <= d <= 31 and 1 <= mo <= 12:
                return f"{d:02d}/{mo:02d}/{year}"
        except ValueError:
            return ""

    # Dernier recours : parsing pandas (gère serials Excel déjà convertis
    # en str, formats US, etc.)
    try:
        dt = pd.to_datetime(str(value).strip(), dayfirst=True, errors="raise")
        return dt.strftime("%d/%m/%Y")
    except Exception:
        return ""


# Mots-clés de détection (listes ordonnées = priorité)
DATE_KEYWORDS = [
    "date de transaction",
    "date transaction",
    "date de vente",
    "date vente",
    "date",
]

VOLUME_KEYWORDS = [
    "volume total",
    "volume",
    "quantite",
    "quantité",
    "qte",
    "quant",
    "litres",
    "liters",
    "litre",
    "liter",
]


# ============================================================
# LECTURE EXCEL ROBUSTE (fichiers convertis depuis CSV)
# ============================================================

# Mots-clés larges servant à repérer la ligne d'en-tête et la
# meilleure feuille quand un .xlsx converti contient des lignes
# de titre, des cellules fusionnées ou plusieurs feuilles.
_HEADER_HINTS = (
    "moyen de paiement", "carte", "card", "pan", "badge", "payment",
    "montant", "amount", "total",
    "truck", "track", "immat", "plaque", "vehicule", "vehicle", "camion",
    "transaction", "station", "date", "ticket",
)


def _header_score(cells):
    """Score une ligne candidate comme en-tête (nb de cellules indices)."""
    score = 0
    for cell in cells:
        norm = _normalize_name(cell)
        if not norm or norm.startswith("unnamed"):
            continue
        for hint in _HEADER_HINTS:
            if hint in norm:
                score += 1
                break
    return score


def _clean_excel_df(df):
    """Nettoie un DataFrame Excel : colonnes/lignes vides, Unnamed, BOM."""
    # Supprimer lignes et colonnes totalement vides
    df = df.dropna(axis=0, how="all")
    df = df.dropna(axis=1, how="all")

    cleaned_cols = []
    for c in df.columns:
        name = str(c).replace("\ufeff", "").strip()
        # Colonne issue d'une cellule fusionnée vide -> "Unnamed: N" ou "nan"
        if name.lower().startswith("unnamed") or name.lower() == "nan":
            name = ""
        cleaned_cols.append(name)
    df.columns = cleaned_cols

    # Supprimer (par position) les colonnes sans nom ET sans données
    keep_pos = []
    for i, col in enumerate(df.columns):
        if col != "":
            keep_pos.append(i)
            continue
        try:
            series = df.iloc[:, i].dropna().astype(str)
            if any(
                str(v).strip() not in ("", "nan", "NaN", "None")
                for v in series
            ):
                keep_pos.append(i)
        except Exception:
            pass
    if len(keep_pos) != len(df.columns):
        df = df.iloc[:, keep_pos]

    df = df.reset_index(drop=True)
    return df


def _read_excel_smart(raw_bytes, engine):
    """Lit un Excel en gérant les conversions CSV->XLSX.

    - Essaie toutes les feuilles, garde celle qui ressemble le plus
      à un fichier attendu (score sur les en-têtes).
    - Détecte la ligne d'en-tête parmi les 12 premières lignes
      (les exports convertis ont souvent 1-3 lignes de titre).
    """
    try:
        xls = pd.ExcelFile(io.BytesIO(raw_bytes), engine=engine)
    except Exception as e:
        raise ValueError(f"Impossible de lire le fichier Excel. ({e})")

    sheet_names = xls.sheet_names or [0]
    best_df = None
    best_score = -1
    best_sheet = None
    best_header_row = 0

    for sheet in sheet_names:
        try:
            preview = pd.read_excel(
                xls, sheet_name=sheet, header=None, nrows=12, dtype=str
            )
        except Exception:
            continue

        # Chercher la ligne d'en-tête (défaut : ligne 0)
        header_row = 0
        header_best = -1
        max_row = min(len(preview), 12)
        for r in range(max_row):
            try:
                cells = [str(v) for v in preview.iloc[r].tolist()]
            except Exception:
                continue
            # Ignorer les lignes quasi vides
            nonempty = [c for c in cells if c.strip() not in ("", "nan", "NaN", "None")]
            if len(nonempty) < 2:
                continue
            score = _header_score(cells)
            if score > header_best:
                header_best = score
                header_row = r

        try:
            df = pd.read_excel(
                xls, sheet_name=sheet, header=header_row, dtype=str
            )
        except Exception:
            continue

        df = _clean_excel_df(df)
        if df.empty or len(df.columns) == 0:
            continue

        score = _header_score([str(c) for c in df.columns])
        # Bonus : feuille avec des données
        score += min(len(df), 100) / 1000.0
        if score > best_score:
            best_score = score
            best_df = df
            best_sheet = sheet
            best_header_row = header_row

    if best_df is None:
        raise ValueError(
            "Impossible de lire le fichier Excel : "
            "aucune feuille exploitable trouvée."
        )

    # Repérer les colonnes stockées comme NOMBRES Excel (data_type 'n').
    # Une carte de 16-18 chiffres stockée en nombre est arrondie par Excel
    # (15 chiffres significatifs) : la perte est invisible après lecture
    # (ex. ...5003 -> ...5056), d'où un avertissement ciblé.
    try:
        best_df.attrs["excel_numeric_cols"] = _excel_numeric_columns(
            raw_bytes, engine, best_sheet, best_header_row
        )
    except Exception:
        best_df.attrs["excel_numeric_cols"] = set()

    return best_df


def _excel_numeric_columns(raw_bytes, engine, sheet, header_row):
    """Noms normalisés des colonnes majoritairement numériques (xlsx only)."""
    if engine != "openpyxl":
        return set()
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(raw_bytes), read_only=True, data_only=True)
    try:
        ws = wb[sheet] if isinstance(sheet, str) else wb.worksheets[0]
        all_rows = list(ws.iter_rows(values_only=False))
    finally:
        try:
            wb.close()
        except Exception:
            pass
    if header_row >= len(all_rows):
        return set()
    header_cells = all_rows[header_row]
    data_rows = all_rows[header_row + 1:header_row + 201]
    numeric = set()
    for idx, hcell in enumerate(header_cells):
        raw_name = "" if hcell.value is None else str(hcell.value)
        name = raw_name.replace("\ufeff", "").strip()
        if name.lower().startswith("unnamed") or name.lower() == "nan":
            continue
        total = 0
        nums = 0
        for row in data_rows:
            if idx >= len(row):
                continue
            cell = row[idx]
            if cell.value is None or str(cell.value).strip() == "":
                continue
            total += 1
            if cell.data_type == "n":
                nums += 1
        if total > 0 and nums / total >= 0.5:
            numeric.add(_normalize_name(name))
    return numeric


# ============================================================
# LIRE FICHIER
# ============================================================

def read_file(file):
    filename = (getattr(file, "filename", "") or "").lower()

    # Remettre le curseur au début (objet Flask FileStorage)
    try:
        file.seek(0)
    except Exception:
        pass

    if filename.endswith(".xlsx") or filename.endswith(".xls"):
        engine = "openpyxl" if filename.endswith(".xlsx") else "xlrd"
        if filename.endswith(".xls"):
            try:
                import xlrd  # noqa: F401
            except ImportError:
                raise ValueError(
                    "Format .xls non supporté : installez la dépendance "
                    "'xlrd' (pip install xlrd)."
                )
        # Lire les bytes une fois (objet Flask FileStorage + ExcelFile)
        try:
            raw = file.read()
        except Exception:
            raw = b""
        if isinstance(raw, str):
            raw = raw.encode("utf-8")
        if not raw:
            raise ValueError("Le fichier Excel est vide.")
        return _read_excel_smart(raw, engine)

    if filename.endswith(".csv"):
        # Lire les bytes une fois, puis tester plusieurs encodings
        # (exports FR : utf-8-sig, cp1252, latin1, avec séparateur ; ou ,).
        try:
            raw = file.read()
        except Exception:
            raw = b""

        if isinstance(raw, str):
            raw = raw.encode("utf-8")

        if not raw:
            raise ValueError("Le fichier CSV est vide.")

        last_error = None
        for encoding in ("utf-8-sig", "cp1252", "latin1"):
            try:
                text = raw.decode(encoding)
            except (UnicodeDecodeError, LookupError) as e:
                last_error = e
                continue

            # Essai détection auto du séparateur, puis fallbacks FR
            for sep in (None, ";", ",", "\t", "|"):
                try:
                    if sep is None:
                        df = pd.read_csv(
                            io.StringIO(text),
                            dtype=str,
                            sep=None,
                            engine="python",
                            keep_default_na=False,
                            na_values=[],
                        )
                    else:
                        df = pd.read_csv(
                            io.StringIO(text),
                            dtype=str,
                            sep=sep,
                            engine="python",
                            keep_default_na=False,
                            na_values=[],
                        )
                    # Vérifier que le parsing a du sens (>1 col ou header plausible)
                    if df.shape[1] >= 1 and len(df.columns) >= 1:
                        # Si une seule colonne mais le header contient ; ou \t,
                        # c'est probablement un mauvais séparateur -> essayer suivant
                        if df.shape[1] == 1 and any(
                            s in str(df.columns[0]) for s in (";", "\t", "|")
                        ):
                            last_error = ValueError("Mauvais séparateur détecté.")
                            continue
                        return df
                except Exception as e:
                    last_error = e
                    continue

        raise ValueError(
            f"Impossible de lire le fichier CSV. "
            f"Vérifiez le séparateur (;/,) et l'encodage (UTF-8/Windows-1252). "
            f"({last_error})"
        )

    raise ValueError(
        "Format non supporté. Utilisez Excel (.xlsx/.xls) ou CSV."
    )


# ============================================================
# TROUVER COLONNE
# ============================================================

def _normalize_name(text):
    import unicodedata

    text = str(text).replace("\ufeff", "").strip().lower()
    text = text.replace("°", "o").replace("’", "'").replace("`", "'")
    # Retirer les accents : numéro -> numero
    text = "".join(
        c for c in unicodedata.normalize("NFKD", text)
        if not unicodedata.combining(c)
    )
    # Normaliser séparateurs pour la comparaison
    text = re.sub(r"\s+", " ", text)
    return text


def _short_columns(df, limit=12):
    """Liste courte des colonnes détectées pour les messages d'erreur."""
    try:
        cols = [str(c).strip() for c in list(df.columns)[:limit]]
        cols = [c if c != "" else "(sans nom)" for c in cols]
        suffix = ", ..." if len(df.columns) > limit else ""
        return "[" + ", ".join(cols) + "]" + suffix
    except Exception:
        return "[?]"


def find_column(df, keywords):

    normalized_cols = [
        (_normalize_name(c), c) for c in df.columns
    ]

    # Priorité aux mots-clés : le premier mot-clé de la liste gagne
    # (important quand plusieurs colonnes "Montant..." existent).
    for keyword in keywords:
        norm_key = _normalize_name(keyword)
        if not norm_key:
            continue
        for norm_col, original_col in normalized_cols:
            if norm_key in norm_col:
                return original_col

    return None


def _amount_candidates(df, keywords):
    """Toutes les colonnes contenant un des mots-clés (ordre = priorité)."""
    candidates = []
    seen = set()
    for keyword in keywords:
        norm_key = _normalize_name(keyword)
        if not norm_key:
            continue
        for column in df.columns:
            if column in seen:
                continue
            if norm_key in _normalize_name(column):
                seen.add(column)
                candidates.append((keyword, column))
    return candidates


def amount_parse_rates(df, keywords, max_sample=200):
    """Taux de valeurs montants par colonne candidate.

    Retourne [(colonne, nb_lisibles, nb_échantillonnées), ...].
    Utilisé pour le choix ET pour les messages d'erreur explicites.
    """
    rates = []
    for _, column in _amount_candidates(df, keywords):
        try:
            series = df[column].dropna().astype(str)
        except Exception:
            rates.append((column, 0, 0))
            continue
        sample = [v for v in series if str(v).strip() != ""][:max_sample]
        count = 0
        for v in sample:
            number = _parse_amount_number(v)
            if number is not None and number != 0.0:
                count += 1
        rates.append((column, count, len(sample)))
    return rates


def find_best_amount_column(df, keywords):
    """Choisit la meilleure colonne de montant.

    1. Trouve toutes les colonnes contenant un des mots-clés.
    2. Score chacune sur un échantillon : nb de valeurs numériques + somme.
    3. Retourne la meilleure. Si aucune valeur numérique nulle part,
       retombe sur find_column (priorité mots-clés).
    """
    candidates = _amount_candidates(df, keywords)

    if not candidates:
        return None

    if len(candidates) == 1:
        return candidates[0][1]

    best_col = None
    best_score = (-1, -1.0)
    for _, column in candidates:
        series = df[column].dropna().astype(str)
        # Échantillon max 200 valeurs non vides
        sample = [v for v in series if str(v).strip() != ""][:200]
        if not sample:
            score = (0, 0.0)
        else:
            count = 0
            total = 0.0
            for v in sample:
                number = _parse_amount_number(v)
                if number is not None and number != 0.0:
                    count += 1
                    total += abs(number)
            # Bonus si le nom contient "transaction/total/vente/ttc" ;
            # bonus plus fort pour la colonne exactement "Montant"/"Amount"
            # (face à "Montant facturé/frais/remise..." quand les deux
            # sont numériques, la colonne brute gagne par défaut)
            norm = _normalize_name(column)
            bonus = 0.0
            if norm in ("montant", "amount", "montant total", "montant transaction"):
                bonus = 1.0
            elif any(k in norm for k in ("transaction", "total", "ttc", "vente", "paye")):
                bonus = 0.5
            score = (count + bonus, total)
        if score > best_score:
            best_score = score
            best_col = column

    if best_col is None or best_score[0] <= 0:
        # Aucune colonne ne contient de montants -> priorité mots-clés
        return find_column(df, keywords)

    return best_col


# Libellés de lignes récapitulatives (pied de page Excel/CSV) à ignorer :
# leurs montants fausseraient le choix de la colonne et les totaux.
_PSEUDO_CARDS = frozenset({
    "total", "totaux", "total general", "sous-total", "sous total",
    "subtotal", "resume", "sommaire",
})


def _is_pseudo_card(card):
    if card == "":
        return True
    if card[:1].isdigit():
        return False
    norm = _normalize_name(card)
    return norm in _PSEUDO_CARDS or norm.startswith("total ")


# ============================================================
# TRAITEMENT PRINCIPAL
# ============================================================


def process_files(transaction_file, truck_file):

    # ========================================================
    # FICHIER 1 : TRANSACTIONS
    # ========================================================

    transactions = read_file(transaction_file)

    transactions.columns = [
        str(c).strip()
        for c in transactions.columns
    ]

    # Colonne carte (ordre = priorité : le plus spécifique d'abord ;
    # le fichier contient plusieurs colonnes "Numéro de..." : ticket,
    # station, client... seule la colonne carte doit matcher)
    card_column = find_column(
        transactions,
        [
            "numero du moyen de paiement",
            "numero du moyen de",
            "moyen de paiement",
            "du moyen",
            "numero de la carte",
            "de la carte",
            "numero de carte",
            "numero carte",
            "no carte",
            "num carte",
            "n carte",
            "no de carte",
            "num de carte",
            "carte",
            "card",
            "pan",
            "badge",
            "payment"
        ]
    )

    AMOUNT_KEYWORDS = [
        "montant transaction",
        "montant total",
        "montant ttc",
        "montant vente",
        "montant règlement",
        "montant reglement",
        "montant payé",
        "montant paye",
        "montant",
        "total transaction",
        "total",
        "amount"
    ]

    if card_column is None:
        raise ValueError(
            "Colonne 'Numéro du moyen de paiement' "
            "introuvable dans le fichier 1. "
            f"Colonnes détectées : {_short_columns(transactions)}. "
            "Si votre .xlsx converti a des lignes de titre au-dessus, "
            "elles sont ignorées automatiquement : vérifiez que les noms "
            "de colonnes sont présents."
        )

    # Nettoyage carte (+ détection de troncation Excel >15 chiffres) :
    # 1) marqueurs flottants visibles (E+, .0) ; 2) colonne stockée en
    # nombre dans un .xlsx avec des cartes de 16+ chiffres (perte
    # invisible : ...5003 -> ...5056).
    def _card_risky(raw):
        if card_precision_risk(raw):
            return True
        try:
            numeric_cols = transactions.attrs.get("excel_numeric_cols", set())
        except Exception:
            numeric_cols = set()
        if _normalize_name(card_column) in numeric_cols:
            text = "" if pd.isna(raw) else str(raw).strip()
            if re.fullmatch(r"\d{16,}", text):
                return True
        return False

    try:
        risk_count = int(transactions[card_column].apply(_card_risky).sum())
    except Exception:
        risk_count = 0

    transactions["Carte"] = (
        transactions[card_column]
        .apply(clean_card)
    )

    # Supprimer cartes vides + lignes récapitulatives (TOTAL...) AVANT
    # le choix de la colonne montant : leurs valeurs fausseraient le vote
    # et feraient passer le garde-fou à tort.
    transactions = transactions[
        ~transactions["Carte"].apply(_is_pseudo_card)
    ].copy()

    if transactions.empty:
        raise ValueError(
            "Aucune transaction valide trouvée."
        )

    # Colonne montant : votée UNIQUEMENT sur les lignes à carte valide.
    # (Évite de prendre "Montant prix" / "Solde" quand le fichier
    # contient plusieurs colonnes Montant, ou une ligne TOTAL en pied.)
    amount_column = find_best_amount_column(transactions, AMOUNT_KEYWORDS)

    if amount_column is None:
        raise ValueError(
            "Colonne 'Montant' "
            "introuvable dans le fichier 1. "
            f"Colonnes détectées : {_short_columns(transactions)}."
        )

    # Exemples bruts AVANT nettoyage (pour un message d'erreur utile ;
    # attention : si amount_column == "Montant", l'assignation ci-dessous
    # écraserait la colonne source, donc on capture avant).
    raw_amount_samples = (
        transactions[amount_column]
        .dropna().astype(str).str.strip()
    )
    raw_amount_samples = [s for s in raw_amount_samples if s != ""][:5]

    # Taux de lecture par colonne candidate (diagnostic).
    amount_rates = amount_parse_rates(transactions, AMOUNT_KEYWORDS)
    rates_text = ", ".join(
        f"'{col}': {ok}/{n}" for col, ok, n in amount_rates
    )

    # Nettoyage montant
    transactions["Montant"] = (
        transactions[amount_column]
        .apply(clean_amount)
    )

    # Garde-fou : si aucun montant n'est lisible, aider l'utilisateur
    # au lieu d'afficher silencieusement des totaux à 0.
    try:
        nonzero = int((transactions["Montant"] != 0.0).sum())
    except Exception:
        nonzero = 0
    if nonzero == 0 and len(transactions) > 0:
        raise ValueError(
            f"Aucun montant lisible dans la colonne '{amount_column}'. "
            f"Exemples de valeurs : {raw_amount_samples}. "
            f"Taux par colonne : {rates_text}. "
            f"Vérifiez que la colonne du montant de transaction est bien "
            f"présente (pas 'prix unitaire' ni 'solde')."
        )

    # ========================================================
    # COLONNES DATE + VOLUME (optionnelles)
    # ========================================================

    date_column = find_column(transactions, DATE_KEYWORDS)
    volume_column = find_column(transactions, VOLUME_KEYWORDS)

    if date_column is not None:
        transactions["Date"] = (
            transactions[date_column]
            .apply(clean_date)
        )
    else:
        transactions["Date"] = ""

    if volume_column is not None:
        transactions["Volume"] = (
            transactions[volume_column]
            .apply(clean_amount)
        )
    else:
        transactions["Volume"] = 0.0

    # ========================================================
    # ========================================================
    # CALCUL DES TRANSACTIONS (par carte + date)
    # ========================================================

    # Résumé par carte + date (une ligne par carte et par jour ;
    # si pas de colonne date, Date = "" donc une ligne par carte)
    summary = (
        transactions
        .groupby(["Carte", "Date"], as_index=False, dropna=False)
        .agg(
            Nb_transactions=("Carte", "size"),
            Montant_total=("Montant", "sum"),
            Volume_total=("Volume", "sum"),
        )
    )
    summary["Date"] = summary["Date"].fillna("")

    # ========================================================
    # FICHIER 2 : NB TRUCK
    # ========================================================

    trucks = read_file(truck_file)

    trucks.columns = [
        str(c).strip()
        for c in trucks.columns
    ]

    # Colonne carte
    truck_card_column = find_column(
        trucks,
        [
            "numero du moyen de paiement",
            "numero du moyen de",
            "moyen de paiement",
            "du moyen",
            "numero de la carte",
            "de la carte",
            "numero de carte",
            "numero carte",
            "no carte",
            "num carte",
            "n carte",
            "no de carte",
            "num de carte",
            "carte",
            "card",
            "pan",
            "badge",
            "payment"
        ]
    )

    # Colonne nb truck
    truck_value_column = find_column(
        trucks,
        [
            "nb truck",
            "nb track",
            "nb trak",
            "truck",
            "track",
            "immat",
            "plaque",
            "vehicule",
            "vehicle",
            "camion"
        ]
    )

    if truck_card_column is None:
        raise ValueError(
            "Colonne 'Numéro du moyen de paiement' "
            "introuvable dans le fichier 2. "
            f"Colonnes détectées : {_short_columns(trucks)}."
        )

    if truck_value_column is None:
        raise ValueError(
            "Colonne 'nb truck' "
            "introuvable dans le fichier 2. "
            f"Colonnes détectées : {_short_columns(trucks)}."
        )

    # Nettoyage carte
    trucks["Carte"] = (
        trucks[truck_card_column]
        .apply(clean_card)
    )

    # Nettoyage nb truck
    trucks["nb truck"] = (
        trucks[truck_value_column]
        .fillna("")
        .astype(str)
        .str.strip()
    )

    # Garder uniquement les colonnes nécessaires
    trucks = trucks[
        ["Carte", "nb truck"]
    ]

    # Supprimer cartes vides
    trucks = trucks[
        trucks["Carte"] != ""
    ]

    # Une seule ligne par carte
    trucks = trucks.drop_duplicates(
        subset=["Carte"],
        keep="first"
    )

    # ========================================================
    # MERGE
    # ========================================================

    result = pd.merge(
        summary,
        trucks,
        on="Carte",
        how="left"
    )

    result["nb truck"] = (
        result["nb truck"]
        .fillna("")
    )

    # Réorganiser les colonnes : Carte | nb truck | Date | Nb | Volume | Montant
    result = result[
        [
            "Carte",
            "nb truck",
            "Date",
            "Nb_transactions",
            "Volume_total",
            "Montant_total"
        ]
    ]

    # Trier par date chronologique puis carte
    # (Date = JJ/MM/AAAA : tri lexical faux sur plusieurs mois,
    # donc tri sur datetime converti)
    result["_sort_date"] = pd.to_datetime(
        result["Date"], format="%d/%m/%Y", errors="coerce"
    )
    result = result.sort_values(
        by=["_sort_date", "Date", "Carte"]
    ).drop(columns=["_sort_date"]).reset_index(drop=True)

    # ========================================================
    # DETAIL (transactions + truck) pour la feuille Transactions
    # ========================================================

    detailed = pd.merge(
        transactions[["Carte", "Date", "Volume", "Montant"]],
        trucks,
        on="Carte",
        how="left"
    )
    detailed["nb truck"] = detailed["nb truck"].fillna("")
    detailed = detailed[
        ["Carte", "nb truck", "Date", "Volume", "Montant"]
    ]
    detailed["_sort_date"] = pd.to_datetime(
        detailed["Date"], format="%d/%m/%Y", errors="coerce"
    )
    detailed = detailed.sort_values(
        by=["_sort_date", "Date", "Carte"]
    ).drop(columns=["_sort_date"]).reset_index(drop=True)

    # ========================================================
    # TOTAUX PAR DATE + GROUPES POUR SOUS-TOTAUX (tableau Détail)
    # ========================================================

    detailed["Date"] = detailed["Date"].fillna("")
    if detailed.empty:
        per_date = pd.DataFrame(
            columns=["Date", "Nb_transactions", "Volume_total", "Montant_total"]
        )
        detail_groups = []
    else:
        per_date = (
            detailed
            .groupby("Date", as_index=False, dropna=False)
            .agg(
                Nb_transactions=("Carte", "size"),
                Volume_total=("Volume", "sum"),
                Montant_total=("Montant", "sum"),
            )
        )
        per_date["Date"] = per_date["Date"].fillna("")
        per_date["_sort_date"] = pd.to_datetime(
            per_date["Date"], format="%d/%m/%Y", errors="coerce"
        )
        per_date = per_date.sort_values(
            by=["_sort_date", "Date"]
        ).drop(columns=["_sort_date"]).reset_index(drop=True)

        detail_groups = []
        for _, prow in per_date.iterrows():
            date_label = prow["Date"]
            subset = detailed[detailed["Date"] == date_label]
            detail_groups.append({
                "date": date_label if date_label != "" else "(Sans date)",
                "date_key": date_label,
                "rows": subset.to_dict(orient="records"),
                "nb": int(prow["Nb_transactions"]),
                "volume": float(prow["Volume_total"]),
                "montant": float(prow["Montant_total"]),
            })

    # ========================================================
    # MOIS DISTINCTS (MM/AAAA) POUR LE FILTRE MOIS DU DÉTAIL
    # ========================================================

    if detailed.empty:
        detail_months = []
        empty_date_count = 0
    else:
        _months = (
            detailed["Date"]
            .str.extract(r"^\d{2}/(\d{2})/(\d{4})$")
            .dropna()
        )
        if _months.empty:
            detail_months = []
        else:
            _months["month"] = _months[0] + "/" + _months[1]
            _counts = _months["month"].value_counts()
            _order = sorted(
                _counts.index.tolist(),
                key=lambda m: (int(m[3:]), int(m[:2])),
            )
            detail_months = [
                {"month": m, "nb": int(_counts[m])} for m in _order
            ]
        empty_date_count = int((detailed["Date"] == "").sum())

    # ========================================================
    # CARTES NON TROUVÉES (dans CSV mais absentes du mapping)
    # ========================================================

    matched_cards = set(
        trucks.loc[trucks["Carte"] != "", "Carte"].tolist()
    )
    unmatched_base = result[~result["Carte"].isin(matched_cards)].copy()
    unmatched = (
        unmatched_base
        .groupby("Carte", as_index=False)
        .agg(
            Nb_transactions=("Nb_transactions", "sum"),
            Volume_total=("Volume_total", "sum"),
            Montant_total=("Montant_total", "sum"),
        )
        .sort_values(by="Carte")
        .reset_index(drop=True)
    )

    # ========================================================
    # TOTAUX PAR CAMION
    # ========================================================

    with_truck = result[result["nb truck"] != ""].copy()
    if with_truck.empty:
        per_truck = pd.DataFrame(
            columns=[
                "nb truck",
                "Nombre_cartes",
                "Nombre_transactions",
                "Volume_total",
                "Montant_total",
            ]
        )
    else:
        per_truck = (
            with_truck
            .groupby("nb truck", as_index=False)
            .agg(
                Nombre_cartes=("Carte", "nunique"),
                Nombre_transactions=("Nb_transactions", "sum"),
                Volume_total=("Volume_total", "sum"),
                Montant_total=("Montant_total", "sum"),
            )
            .sort_values(by="nb truck")
            .reset_index(drop=True)
        )

    # ========================================================
    # TOTAUX
    # ========================================================

    total_transactions = int(
        result["Nb_transactions"].sum()
    )

    total_amount = float(
        result["Montant_total"].sum()
    )

    total_volume = float(
        result["Volume_total"].sum()
    )

    total_cards = int(result["Carte"].nunique())

    unmatched_count = int(unmatched["Carte"].nunique()) if not unmatched.empty else 0

    card_warning = ""
    if risk_count > 0:
        card_warning = (
            f"Attention : {risk_count} numéro(s) de carte semblent lus "
            "comme des nombres Excel (précision limitée à 15 chiffres, "
            "derniers chiffres possiblement arrondis). Pour des numéros "
            "exacts, importez le CSV d'origine ou formatez la colonne "
            "carte en Texte avant la conversion en .xlsx."
        )

    # Diagnostic : colonnes détectées + taux de lecture (affiché sous
    # le résultat et dans la console serveur pour faciliter le support).
    diag_parts = [
        f"Carte: '{card_column}'",
        f"Montant: '{amount_column}' ({nonzero}/{len(transactions)} valeurs)",
        f"Volume: '{volume_column}'" if volume_column else "Volume: (absent -> 0)",
        f"Date: '{date_column}'" if date_column else "Date: (absente)",
    ]
    diagnostics = "Détection — " + " | ".join(diag_parts)
    print(f"[PaymentAnalyzer] {diagnostics} ; taux: {rates_text}")

    return {
        "result": result,
        "detailed": detailed,
        "detail_groups": detail_groups,
        "detail_months": detail_months,
        "empty_date_count": empty_date_count,
        "per_date": per_date,
        "unmatched": unmatched,
        "per_truck": per_truck,
        "total_transactions": total_transactions,
        "total_amount": total_amount,
        "total_volume": total_volume,
        "total_cards": total_cards,
        "unmatched_count": unmatched_count,
        "card_warning": card_warning,
        "diagnostics": diagnostics,
    }


# ============================================================
# FORMAT MONTANT
# ============================================================

def format_money(value):

    return (
        f"{float(value):,.2f}"
        .replace(",", " ")
        .replace(".", ",")
    )


# ============================================================
# PAGE PRINCIPALE
# ============================================================

# ============================================================
# EXPORT EXCEL 5 FEUILLES
# ============================================================

def _with_total(df, total_values):
    total_row = pd.DataFrame([total_values])
    return pd.concat([df, total_row], ignore_index=True)


def _style_sheet(writer, sheet_name, widths, money_cols):
    worksheet = writer.sheets[sheet_name]
    for letter, width in widths.items():
        worksheet.column_dimensions[letter].width = width
    max_row = worksheet.max_row
    for col in money_cols:
        for cell in worksheet[col][1:max_row]:
            cell.number_format = '#,##0.00'
    # Gras sur les lignes SOUS-TOTAL / TOTAL (col A)
    try:
        from openpyxl.styles import Font, PatternFill
        bold = Font(bold=True)
        sub_fill = PatternFill(start_color="E8EEF4", fill_type="solid")
        total_fill = PatternFill(start_color="EEEEEE", fill_type="solid")
        for row in worksheet.iter_rows(min_row=2, max_row=max_row):
            first = "" if row[0].value is None else str(row[0].value)
            if first.startswith("SOUS-TOTAL"):
                for cell in row:
                    cell.font = bold
                    cell.fill = sub_fill
            elif first == "TOTAL":
                for cell in row:
                    cell.font = bold
                    cell.fill = total_fill
    except Exception:
        pass


def build_export(data):
    """Construit le classeur 5 feuilles et retourne les bytes."""
    result = data["result"]
    detailed = data["detailed"]
    unmatched = data["unmatched"]
    per_truck = data["per_truck"]

    # Synthèse par carte (sans Date) : Carte | nb truck | Nb | Montant
    if result.empty:
        per_card = pd.DataFrame(
            columns=["Carte", "nb truck", "Nb_transactions", "Montant_total"]
        )
    else:
        per_card = (
            result
            .groupby(["Carte", "nb truck"], as_index=False, dropna=False)
            .agg(
                Nb_transactions=("Nb_transactions", "sum"),
                Montant_total=("Montant_total", "sum"),
            )
            .sort_values(by="Carte")
            .reset_index(drop=True)
        )
    per_card_sheet = _with_total(per_card, {
        "Carte": "TOTAL",
        "nb truck": "",
        "Nb_transactions": int(per_card["Nb_transactions"].sum()) if not per_card.empty else 0,
        "Montant_total": float(per_card["Montant_total"].sum()) if not per_card.empty else 0.0,
    })

    resume = _with_total(result, {
        "Carte": "TOTAL",
        "nb truck": "",
        "Date": "",
        "Nb_transactions": data["total_transactions"],
        "Volume_total": data["total_volume"],
        "Montant_total": data["total_amount"],
    })

    trans_rows = []
    if detailed.empty:
        trans = pd.DataFrame(
            columns=["Carte", "nb truck", "Date", "Volume", "Montant"]
        )
        trans = _with_total(trans, {
            "Carte": "TOTAL",
            "nb truck": "",
            "Date": "",
            "Volume": 0.0,
            "Montant": 0.0,
        })
    else:
        _trans_detail = detailed[
            ["Carte", "nb truck", "Date", "Volume", "Montant"]
        ].copy()
        _trans_detail["_sort_date"] = pd.to_datetime(
            _trans_detail["Date"], format="%d/%m/%Y", errors="coerce"
        )
        _trans_detail = _trans_detail.sort_values(
            by=["_sort_date", "Date", "Carte"]
        ).drop(columns=["_sort_date"]).reset_index(drop=True)
        for date_label, subset in _trans_detail.groupby("Date", sort=False):
            trans_rows.append(subset[["Carte", "nb truck", "Date", "Volume", "Montant"]])
            date_txt = date_label if str(date_label) != "" else "(Sans date)"
            trans_rows.append(pd.DataFrame([{
                "Carte": f"SOUS-TOTAL {date_txt} ({len(subset)})",
                "nb truck": "",
                "Date": date_label,
                "Volume": float(subset["Volume"].sum()),
                "Montant": float(subset["Montant"].sum()),
            }]))
        trans = pd.concat(trans_rows, ignore_index=True) if trans_rows else _trans_detail[
            ["Carte", "nb truck", "Date", "Volume", "Montant"]
        ]
        trans = _with_total(trans, {
            "Carte": "TOTAL",
            "nb truck": "",
            "Date": "",
            "Volume": float(detailed["Volume"].sum()),
            "Montant": float(detailed["Montant"].sum()),
        })

    if unmatched.empty:
        notfound = pd.DataFrame(
            columns=["Carte", "Nb_transactions", "Volume_total", "Montant_total"]
        )
    else:
        notfound = _with_total(unmatched, {
            "Carte": "TOTAL",
            "Nb_transactions": int(unmatched["Nb_transactions"].sum()),
            "Volume_total": float(unmatched["Volume_total"].sum()),
            "Montant_total": float(unmatched["Montant_total"].sum()),
        })

    if per_truck.empty:
        trucks_sheet = pd.DataFrame(
            columns=[
                "nb truck", "Nombre_cartes", "Nombre_transactions",
                "Volume_total", "Montant_total",
            ]
        )
    else:
        trucks_sheet = _with_total(per_truck, {
            "nb truck": "TOTAL",
            "Nombre_cartes": int(per_truck["Nombre_cartes"].sum()),
            "Nombre_transactions": int(per_truck["Nombre_transactions"].sum()),
            "Volume_total": float(per_truck["Volume_total"].sum()),
            "Montant_total": float(per_truck["Montant_total"].sum()),
        })

    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        resume.to_excel(writer, index=False, sheet_name="Résumé")
        _style_sheet(
            writer, "Résumé",
            {"A": 22, "B": 16, "C": 14, "D": 18, "E": 18, "F": 18},
            ["E", "F"],
        )

        trans.to_excel(writer, index=False, sheet_name="Transactions")
        _style_sheet(
            writer, "Transactions",
            {"A": 22, "B": 16, "C": 14, "D": 16, "E": 16},
            ["D", "E"],
        )

        notfound.to_excel(writer, index=False, sheet_name="Cartes non trouvées")
        _style_sheet(
            writer, "Cartes non trouvées",
            {"A": 22, "B": 20, "C": 18, "D": 18},
            ["C", "D"],
        )

        trucks_sheet.to_excel(writer, index=False, sheet_name="Totaux par camion")
        _style_sheet(
            writer, "Totaux par camion",
            {"A": 18, "B": 18, "C": 22, "D": 18, "E": 18},
            ["D", "E"],
        )

        per_card_sheet.to_excel(writer, index=False, sheet_name="Synthèse par carte")
        _style_sheet(
            writer, "Synthèse par carte",
            {"A": 22, "B": 16, "C": 20, "D": 18},
            ["D"],
        )

    output.seek(0)
    return output.getvalue()


@app.route("/", methods=["GET", "POST"])
def index():

    result = None

    detailed = None

    detail_groups = []

    detail_months = []

    empty_date_count = 0

    per_date = None

    total_transactions = 0

    total_amount = 0

    total_volume = 0.0

    total_cards = 0

    unmatched_count = 0

    card_warning = ""

    diagnostics = ""

    error = None

    download_url = None

    if request.method == "POST":

        transaction_file = request.files.get(
            "transaction_file"
        )

        truck_file = request.files.get(
            "truck_file"
        )

        # ====================================================
        # VERIFICATION FICHIERS
        # ====================================================

        if not transaction_file:

            error = (
                "Veuillez sélectionner le fichier transactions."
            )

        elif not truck_file:

            error = (
                "Veuillez sélectionner le fichier nb truck."
            )

        elif transaction_file.filename == "":

            error = (
                "Le fichier transactions est vide."
            )

        elif truck_file.filename == "":

            error = (
                "Le fichier nb truck est vide."
            )

        else:

            try:

                # =================================================
                # TRAITEMENT
                # =================================================

                data = process_files(
                    transaction_file,
                    truck_file
                )

                result = data["result"]
                detailed = data["detailed"]
                detail_groups = data.get("detail_groups", [])
                detail_months = data.get("detail_months", [])
                empty_date_count = data.get("empty_date_count", 0)
                per_date = data.get("per_date")
                total_transactions = data["total_transactions"]
                total_amount = data["total_amount"]
                total_volume = data["total_volume"]
                total_cards = data["total_cards"]
                unmatched_count = data["unmatched_count"]
                card_warning = data.get("card_warning", "")
                diagnostics = data.get("diagnostics", "")

                # =================================================
                # CREATION EXCEL (5 feuilles)
                # =================================================

                excel_bytes = build_export(data)

                # =================================================
                # SAUVEGARDER LE FICHIER SUR LE SERVEUR
                # =================================================

                file_id = str(uuid.uuid4())

                file_path = os.path.join(
                    GENERATED_FOLDER,
                    f"{file_id}.xlsx"
                )

                with open(
                    file_path,
                    "wb"
                ) as f:

                    f.write(excel_bytes)

                # =================================================
                # URL DE TELECHARGEMENT
                # =================================================

                download_url = url_for(
                    "export",
                    file_id=file_id
                )

            except Exception as e:

                error = str(e)

    return render_template(
        "index.html",

        result=result,

        detailed=detailed,

        detail_groups=detail_groups,

        detail_months=detail_months,

        empty_date_count=empty_date_count,

        per_date=per_date,

        total_transactions=total_transactions,

        total_amount=total_amount,

        total_volume=total_volume,

        total_cards=total_cards,

        unmatched_count=unmatched_count,

        card_warning=card_warning,

        diagnostics=diagnostics,

        format_money=format_money,

        error=error,

        download_url=download_url
    )


# ============================================================
# TELECHARGEMENT EXCEL
# ============================================================

@app.route("/export/<file_id>")
def export(file_id):

    # The generated folder is absolute and works with the EXE.
    file_path = os.path.join(
        GENERATED_FOLDER,
        f"{file_id}.xlsx"
    )

    # Security: only allow the generated UUID-style filename.
    if not re.fullmatch(
        r"[0-9a-fA-F-]{36}",
        file_id
    ):
        return "Identifiant de fichier invalide.", 400

    if not os.path.isfile(file_path):
        return (
            "Fichier Excel introuvable.",
            404
        )

    try:
        return send_file(
            file_path,
            as_attachment=True,
            download_name="resultat_transactions.xlsx",
            mimetype=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            )
        )

    except Exception as e:
        return (
            f"Erreur lors du téléchargement : {str(e)}",
            500
        )


# ============================================================
# START
# ============================================================

if __name__ == "__main__":

    # Open the browser automatically only after the Flask server
    # has had a moment to start.
    def open_browser():
        webbrowser.open("http://127.0.0.1:5000")

    threading.Timer(
        1.5,
        open_browser
    ).start()

    app.run(
        host="127.0.0.1",
        port=5000,
        debug=False,
        use_reloader=False
    )
